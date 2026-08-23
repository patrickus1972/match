"""AT-1 .. AT-8, the acceptance tests in section 6 of the requirements.

These cover everything in the acceptance tests that a script can check. The
conversational halves — "asks at most one clarifying question", "volunteers
before calculating" — are governed by SKILL.md and are checked by hand against
the transcript; the figures, artefacts and refusals they depend on are here.
"""

import json

import pytest
from pptx import Presentation

import dataset
import render_deck
import run_case
from conftest import DATA, load_fixture


def titles(deck_path) -> list[str]:
    found = []
    for slide in Presentation(str(deck_path)).slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                found.append(shape.text_frame.text)
                break
    return found


def slide_text(deck_path, index: int) -> str:
    slide = list(Presentation(str(deck_path)).slides)[index]
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        if getattr(shape, "has_table", False) and shape.has_table:
            parts += [c.text for r in shape.table.rows for c in r.cells]
    return " ".join(parts)


def test_at1_belgian_display_case(out_dir):
    """Resolved inputs, a five-slide deck and a calculation file."""
    record, paths = run_case.run(load_fixture("at1_belgium_display"), out_dir, DATA)
    assert record["result"]["status"] == "OK"
    # INT-2: nothing that is on file had to be asked for.
    for field in ("net_price_eur", "contribution_margin_pct",
                  "base_rate_of_sale_units_week", "seasonality", "category",
                  "need_state", "market"):
        assert record["inputs"][field] is not None
    assert record["result"]["comparables_used"] >= 4
    body = titles(paths["deck"])[:-1]
    assert len(body) == 5
    assert paths["calculation"].exists() and paths["comparables"].exists()


def test_at2_correcting_the_margin_changes_the_figure(out_dir):
    request = load_fixture("at1_belgium_display")
    before, first = run_case.run(request, out_dir / "before", DATA)
    corrected = request | {"contribution_margin_pct": 0.40}
    after, second = run_case.run(corrected, out_dir / "after", DATA)

    original_cm = before["inputs"]["contribution_margin_pct"]
    expected = (before["result"]["base_case"]["gross_profit_eur"]
                * (0.40 / original_cm))
    assert after["result"]["base_case"]["gross_profit_eur"] == pytest.approx(expected)
    assert second["deck"].read_bytes() != first["deck"].read_bytes()
    assert "40.0%" in slide_text(second["deck"], 4)


def test_at3_no_comparables_means_no_number(out_dir):
    record, paths = run_case.run(load_fixture("dat6_no_comparables"), out_dir, DATA)
    assert record["result"]["status"] == "INSUFFICIENT_COMPARABLES"
    assert "scenarios" not in record["result"]
    assert record["assumptions"] == []
    assert titles(paths["deck"]) == ["The ask", render_deck.APPENDIX_TITLE]
    first = slide_text(paths["deck"], 0)
    assert render_deck.NOT_SUPPLIED in first
    assert "Gross profit (base case)" not in first
    # No forecast figure anywhere: the only euro amount is the A&P request.
    blob = json.dumps(record)
    assert "gross_profit_eur" not in blob


def test_at4_cannibalisation_flips_the_case(out_dir):
    record, paths = run_case.run(load_fixture("dat7_cannibalisation"), out_dir, DATA)
    base = record["result"]["base_case"]
    assert base["gross_profit_eur"] < 0
    assert "CANNIBALISATION_DOMINATES" in [a["code"] for a in record["advisories"]]
    portfolio = slide_text(paths["deck"], titles(paths["deck"]).index("Portfolio impact"))
    assert "-" in portfolio          # net units are shown negative
    assert "CANNIBALISED UNITS" in portfolio.upper()


def test_at5_low_below_break_even_is_volunteered_and_shown_first(out_dir):
    record, paths = run_case.run(load_fixture("dat8_low_below_breakeven"), out_dir, DATA)
    assert "LOW_BELOW_BREAK_EVEN" in [a["code"] for a in record["advisories"]]
    sensitivity = slide_text(paths["deck"],
                             titles(paths["deck"]).index("Sensitivity"))
    assert sensitivity.index("Low") < sensitivity.index("Base") < sensitivity.index("High")
    ask = slide_text(paths["deck"], 0)
    assert "does not repay" in ask


@pytest.mark.parametrize("ask", [
    "Should I approve this?",
    "Should we fund it?",
    "Do you recommend this case?",
    "Which should I prioritise, this or the display?",
])
def test_at6_a_request_for_a_recommendation_is_declined(ask):
    assert dataset.is_recommendation_request(ask) is True
    reply = dataset.DECLINE_TEXT.lower()
    for verdict in ("i recommend", "approve it", "worth funding", "go ahead",
                    "i would approve", "strong case"):
        assert verdict not in reply


def test_at6_an_ordinary_question_is_not_treated_as_a_recommendation_ask():
    for ordinary in ("What is the break-even uplift?",
                     "Can you run it at four weeks?",
                     "Where did the 17.4% come from?"):
        assert dataset.is_recommendation_request(ordinary) is False


def test_at7_a_polish_npi_routes_to_npi_rough(out_dir):
    record, paths = run_case.run(load_fixture("at7_npi_rough_poland"), out_dir, DATA)
    assert record["depth"] == "npi_rough"
    assert record["result"]["status"] == "OK"
    body = titles(paths["deck"])[:-1]
    assert len(body) == 4
    assert "payback_months" not in record["result"]


def test_at8_re_rendering_from_the_calculation_file_is_byte_identical(out_dir):
    _, paths = run_case.run(load_fixture("at1_belgium_display"), out_dir, DATA)
    original = paths["deck"].read_bytes()
    paths["deck"].unlink()
    record = json.loads(paths["calculation"].read_text(encoding="utf-8"))
    again = render_deck.render(record, paths["deck"])
    assert again.read_bytes() == original
