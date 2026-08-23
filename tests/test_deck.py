"""OUT-1 .. OUT-11 and AT-8: the rendered deck."""

import json
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

import dataset
import render_deck
import run_case
from conftest import DATA, load_fixture


def build(fixture: str, out_dir):
    return run_case.run(load_fixture(fixture), out_dir, DATA)


def texts(slide) -> list[str]:
    found = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            found.append(shape.text_frame.text)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    found.append(cell.text)
    return found


def titles(deck_path) -> list[str]:
    return [texts(slide)[0] for slide in Presentation(str(deck_path)).slides]


def test_out1_a_run_produces_a_deck_and_a_calculation_file(out_dir):
    record, paths = build("at1_belgium_display", out_dir)
    assert paths["deck"].exists() and paths["deck"].suffix == ".pptx"
    assert paths["calculation"].exists() and paths["calculation"].suffix == ".json"
    assert paths["comparables"].exists() and paths["comparables"].suffix == ".csv"
    assert set(p.name for p in out_dir.iterdir()) == {p.name for p in paths.values()}
    reloaded = json.loads(paths["calculation"].read_text(encoding="utf-8"))
    assert reloaded["result"]["status"] == "OK"


def test_out2_and_at8_rendering_again_from_the_file_is_byte_identical(out_dir):
    _, paths = build("at1_belgium_display", out_dir)
    original = paths["deck"].read_bytes()
    record = json.loads(paths["calculation"].read_text(encoding="utf-8"))
    again = render_deck.render(record, out_dir / "again.pptx")
    assert again.read_bytes() == original


def test_out3_activation_deck_body_slides(out_dir):
    _, paths = build("at1_belgium_display", out_dir)
    found = titles(paths["deck"])
    assert found[:5] == ["The ask", "Forecast", "Portfolio impact", "Sensitivity",
                         "Assumptions and check-back"]
    assert len(found) == 6          # five body slides plus the appendix (OUT-5)


def test_out4_npi_deck_lengths(out_dir):
    _, rough = build("at7_npi_rough_poland", out_dir)
    _, full = build("npi_full_germany", out_dir)
    assert len(titles(rough["deck"])) == 4 + 1
    assert len(titles(full["deck"])) == 8 + 1


def test_out5_every_depth_carries_an_appendix(out_dir):
    for fixture in ("at1_belgium_display", "at7_npi_rough_poland",
                    "npi_full_germany", "dat6_no_comparables"):
        _, paths = build(fixture, out_dir / fixture)
        found = titles(paths["deck"])
        assert found[-1] == render_deck.APPENDIX_TITLE, fixture
        appendix = list(Presentation(str(paths["deck"])).slides)[-1]
        joined = " ".join(texts(appendix))
        assert "A&P request" in joined


def test_out6_slide_order_is_identical_across_cases_of_the_same_depth(out_dir):
    _, first = build("at1_belgium_display", out_dir / "a")
    _, second = build("adv2_no_siblings", out_dir / "b")
    assert titles(first["deck"]) == titles(second["deck"])


def test_out7_footer_on_every_slide(out_dir):
    for fixture in ("at1_belgium_display", "npi_full_germany", "dat6_no_comparables"):
        _, paths = build(fixture, out_dir / fixture)
        for index, slide in enumerate(Presentation(str(paths["deck"])).slides):
            assert render_deck.FOOTER_TEXT in texts(slide), f"{fixture} slide {index+1}"


def test_out8_sensitivity_shows_the_low_scenario_first(out_dir):
    _, paths = build("dat8_low_below_breakeven", out_dir)
    slides = list(Presentation(str(paths["deck"])).slides)
    sensitivity = slides[titles(paths["deck"]).index("Sensitivity")]
    table = next(s.table for s in sensitivity.shapes
                 if getattr(s, "has_table", False) and s.has_table)
    assert [table.cell(r, 0).text for r in range(1, 4)] == ["Low", "Base", "High"]


def test_out9_every_assumption_has_a_source_and_an_owner(out_dir):
    for fixture in ("at1_belgium_display", "npi_full_germany"):
        record, paths = build(fixture, out_dir / fixture)
        assert record["assumptions"]
        for row in record["assumptions"]:
            assert row["source"].strip() and row["owner"].strip()
        slides = list(Presentation(str(paths["deck"])).slides)
        slide = slides[titles(paths["deck"]).index("Assumptions and check-back")]
        table = next(s.table for s in slide.shapes
                     if getattr(s, "has_table", False) and s.has_table)
        for r in range(1, len(table.rows)):
            assert table.cell(r, 2).text.strip()
            assert table.cell(r, 3).text.strip()


def test_out10_insufficient_comparables_deck_has_no_forecast(out_dir):
    record, paths = build("dat6_no_comparables", out_dir)
    assert record["result"]["status"] == "INSUFFICIENT_COMPARABLES"
    found = titles(paths["deck"])
    assert found == ["The ask", render_deck.APPENDIX_TITLE]
    slides = list(Presentation(str(paths["deck"])).slides)
    first = " ".join(texts(slides[0]))
    assert render_deck.NOT_SUPPLIED in first
    for forbidden in ("Forecast", "Sensitivity", "Gross profit (base case)"):
        assert forbidden not in " ".join(found)
    assert "€" not in first.replace("€0", "")  or "NOT SUPPLIED" in first


def test_out10_missing_input_deck_names_the_field(out_dir):
    request = load_fixture("at1_belgium_display") | {"ap_request_eur": None,
                                                     "case_id": "MISSING-AP"}
    ds = dataset.Dataset(DATA)
    record = dataset.assemble(ds, request)
    assert record["result"]["status"] == "REQUIRED_NOT_SUPPLIED"
    assert "ap_request_eur" in record["result"]["missing_fields"]
    paths = run_case.write_artefacts(record, out_dir)
    first = " ".join(texts(list(Presentation(str(paths["deck"])).slides)[0]))
    assert render_deck.NOT_SUPPLIED in first and "ap_request_eur" in first


def test_out11_package_is_well_formed(out_dir):
    _, paths = build("npi_full_germany", out_dir)
    with zipfile.ZipFile(paths["deck"]) as archive:
        assert archive.testzip() is None
        assert "[Content_Types].xml" in archive.namelist()
    assert len(Presentation(str(paths["deck"])).slides) == 9


IMPRESS_AVAILABLE = (
    shutil.which("soffice") is not None
    and Path("/usr/lib/libreoffice/share/registry/impress.xcd").exists()
)


@pytest.mark.skipif(not IMPRESS_AVAILABLE,
                    reason="LibreOffice Impress filters are not installed here; "
                           "OUT-11 stays a manual gate in this environment")
def test_out11_renders_to_pdf(out_dir):
    _, paths = build("at1_belgium_display", out_dir)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf",
                    "--outdir", str(out_dir), str(paths["deck"])],
                   check=True, capture_output=True, timeout=180)
    assert (out_dir / paths["deck"].with_suffix(".pdf").name).exists()
